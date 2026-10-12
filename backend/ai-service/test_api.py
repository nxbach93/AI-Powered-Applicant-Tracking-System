"""
Script Kiểm thử Độc lập (Decoupled Testing) cho cả 2 API cốt lõi:
  1. POST /api/v1/parse-cv: Bóc tách thực thể CV từ file (.pdf / .docx).
  2. POST /api/v1/match-cv: So khớp CV với JD (Matching Engine đa trọng số + NLP).

Chỉ sử dụng thư viện chuẩn của Python kết hợp python-docx (đã có trong venv).
Không phụ thuộc thêm thư viện HTTP bên ngoài (như requests).
"""

import io
import json
import sys
import urllib.error
import urllib.request

try:
    import docx
except ImportError:
    docx = None

PARSE_CV_URL = "http://localhost:8000/api/v1/parse-cv"
MATCH_CV_URL = "http://localhost:8000/api/v1/match-cv"


# ============================================================================
# Helpers gửi Request (Pure Standard Library)
# ============================================================================
def create_sample_docx_bytes() -> bytes:
    """Tạo file DOCX mẫu chuẩn trong bộ nhớ bằng python-docx."""
    if docx is None:
        raise RuntimeError("Thư viện 'python-docx' chưa được cài đặt trong môi trường!")

    doc = docx.Document()
    doc.add_heading("NGUYỄN VĂN AN - SENIOR FRONTEND DEVELOPER", level=1)

    p_contact = doc.add_paragraph()
    p_contact.add_run("Email: candidate@example.com | Phone: 0987654321 | Địa chỉ: Hà Nội\n")

    p_edu = doc.add_paragraph()
    p_edu.add_run("HỌC VẤN:\n").bold = True
    p_edu.add_run("Đại học Bách Khoa Hà Nội - Chuyên ngành Công nghệ Thông tin (2018 - 2022)\n")

    p_exp = doc.add_paragraph()
    p_exp.add_run("KINH NGHIỆM LÀM VIỆC:\n").bold = True
    p_exp.add_run("Có 3 năm kinh nghiệm phát triển hệ thống Web với React, TypeScript và Node.js.\n")

    p_skills = doc.add_paragraph()
    p_skills.add_run("KỸ NĂNG CHUYÊN MÔN:\n").bold = True
    p_skills.add_run("Frontend: React, Next.js, TypeScript, Tailwind CSS, Redux\n")
    p_skills.add_run("Backend: Node.js, Express, Python, FastAPI, RESTful API\n")
    p_skills.add_run("Database & DevOps: MySQL, PostgreSQL, Docker, Git, CI/CD\n")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def send_multipart_request(
    url: str,
    file_bytes: bytes | None = None,
    filename: str | None = None,
    field_name: str = "file",
) -> tuple[int, dict]:
    """Gửi HTTP POST multipart/form-data thuần bằng urllib.request."""
    boundary = "----WebKitFormBoundaryATSValidation7MA4YWxkTr"

    if file_bytes is not None and filename is not None:
        header_part = (
            f"--{boundary}\r\n"
            f'Content-Disposition: form-data; name="{field_name}"; filename="{filename}"\r\n'
            f"Content-Type: application/octet-stream\r\n\r\n"
        ).encode("utf-8")
        footer_part = f"\r\n--{boundary}--\r\n".encode("utf-8")
        body = header_part + file_bytes + footer_part
        content_type = f"multipart/form-data; boundary={boundary}"
    else:
        body = b""
        content_type = "multipart/form-data; boundary=" + boundary

    req = urllib.request.Request(
        url=url,
        data=body,
        headers={"Content-Type": content_type},
        method="POST",
    )

    try:
        with urllib.request.urlopen(req) as response:
            res_body = response.read().decode("utf-8")
            return response.status, json.loads(res_body)
    except urllib.error.HTTPError as exc:
        res_body = exc.read().decode("utf-8")
        try:
            parsed_json = json.loads(res_body)
        except Exception:
            parsed_json = {"raw": res_body}
        return exc.code, parsed_json
    except urllib.error.URLError as err:
        return 0, {"error": f"Không thể kết nối đến máy chủ ({err.reason}). Bạn đã khởi chạy Uvicorn chưa?"}


def send_json_request(url: str, json_data: dict) -> tuple[int, dict]:
    """Gửi HTTP POST application/json thuần bằng urllib.request."""
    body = json.dumps(json_data).encode("utf-8")
    req = urllib.request.Request(
        url=url,
        data=body,
        headers={"Content-Type": "application/json"},
        method="POST",
    )

    try:
        with urllib.request.urlopen(req) as response:
            res_body = response.read().decode("utf-8")
            return response.status, json.loads(res_body)
    except urllib.error.HTTPError as exc:
        res_body = exc.read().decode("utf-8")
        try:
            parsed_json = json.loads(res_body)
        except Exception:
            parsed_json = {"raw": res_body}
        return exc.code, parsed_json
    except urllib.error.URLError as err:
        return 0, {"error": f"Không thể kết nối đến máy chủ ({err.reason})."}


# ============================================================================
# Kịch bản kiểm thử (Test Runner)
# ============================================================================
def run_tests():
    print("=" * 75)
    print("🧪 BẮT ĐẦU KIỂM THỬ ĐỘC LẬP TOÀN BỘ AI MICROSERVICE (GIAI ĐOẠN 1 & 2)")
    print("=" * 75)

    # -----------------------------------------------------------------------
    # GIAI ĐOẠN 1: TEST /parse-cv
    # -----------------------------------------------------------------------
    print("\n[TEST 1] Gửi File CV hợp lệ (.docx) tới /api/v1/parse-cv...")
    parsed_cv_result = None
    try:
        docx_data = create_sample_docx_bytes()
        status_code, resp_data = send_multipart_request(
            url=PARSE_CV_URL,
            file_bytes=docx_data,
            filename="sample_cv.docx",
        )
        print(f"  👉 HTTP Status: {status_code}")
        print("  👉 Response Payload:")
        print(json.dumps(resp_data, indent=2, ensure_ascii=False))

        if status_code == 200 and resp_data.get("success") is True:
            parsed_cv_result = resp_data
            contact = resp_data.get("contact_info", {})
            skills = resp_data.get("parsed_skills", [])
            print(f"  ✅ [PASS] Trích xuất Email='{contact.get('email')}', SĐT='{contact.get('phone')}'")
            print(f"  ✅ [PASS] Số kỹ năng nhận diện: {len(skills)}")
        else:
            print(f"  ❌ [FAIL] Mong đợi 200 OK nhưng nhận mã {status_code}")
    except Exception as exc:
        print(f"  ❌ [ERROR] Lỗi khi chạy Test 1: {exc}")

    # -----------------------------------------------------------------------
    print("\n" + "-" * 75)
    print("[TEST 2] Gửi File CV sai định dạng (.txt)...")
    fake_txt_bytes = b"Ho va ten: Nguyen Van B\nKy nang: Java, Spring"
    status_code, resp_data = send_multipart_request(
        url=PARSE_CV_URL,
        file_bytes=fake_txt_bytes,
        filename="invalid_cv.txt",
    )
    print(f"  👉 HTTP Status: {status_code}")
    if status_code == 415 and resp_data.get("success") is False:
        print(f"  ✅ [PASS] Bắt lỗi 415 Unsupported Media Type chính xác!")
    else:
        print(f"  ❌ [FAIL] Mong đợi 415 nhưng nhận mã {status_code}")

    # -----------------------------------------------------------------------
    # GIAI ĐOẠN 2: TEST /match-cv (MATCHING ENGINE)
    # -----------------------------------------------------------------------
    print("\n" + "=" * 75)
    print("🚀 GIAI ĐOẠN 2: KIỂM THỬ MATCHING ENGINE (/api/v1/match-cv)")
    print("=" * 75)

    print("\n[TEST 3] So khớp CV hợp lệ với JD (Kịch bản chuẩn)...")
    cv_skills_to_use = (
        parsed_cv_result.get("parsed_skills", [])
        if parsed_cv_result
        else ["React", "TypeScript", "Node.js", "Docker", "Tailwind CSS", "SQL"]
    )

    match_payload = {
        "cv_data": {
            "raw_text": (
                "Nguyễn Văn An. 3 năm kinh nghiệm phát triển Web Frontend và Fullstack. "
                "Kỹ năng: React, TypeScript, Node.js, Docker, Tailwind CSS, SQL. "
                "Tốt nghiệp Đại học Bách Khoa Hà Nội."
            ),
            "skills": cv_skills_to_use,
            "experience_years": 3,
        },
        "jd_data": {
            "title": "Senior Frontend Developer (React/Next.js)",
            "description": "Tham gia phát triển hệ thống quản lý tuyển dụng nội bộ quy mô lớn.",
            "requirements": "Thành thạo React, Next.js, TypeScript. Tối thiểu 3 năm kinh nghiệm. Ưu tiên có kiến thức về Docker và GraphQL.",
            "required_skills": ["React", "Next.js", "TypeScript", "GraphQL", "Docker"],
            "min_experience_years": 3,
        },
        "weights": {
            "skill_weight": 0.5,
            "exp_weight": 0.3,
            "semantic_weight": 0.2,
        },
    }

    status_code, resp_data = send_json_request(MATCH_CV_URL, match_payload)
    print(f"  👉 HTTP Status: {status_code}")
    print("  👉 Response Payload:")
    print(json.dumps(resp_data, indent=2, ensure_ascii=False))

    if status_code == 200 and resp_data.get("success") is True:
        match_score = resp_data.get("match_score")
        breakdown = resp_data.get("score_breakdown", {})
        matched = resp_data.get("matched_skills", [])
        missing = resp_data.get("missing_skills", [])
        summary = resp_data.get("summary", "")

        print(f"\n  🎯 [PASS] Điểm tổng hợp Match Score: {match_score:.1f} / 100")
        print(f"  📊 [PASS] Chi tiết điểm: Kỹ năng={breakdown.get('skill_score')}%, Kinh nghiệm={breakdown.get('experience_score')}%, Ngữ nghĩa={breakdown.get('semantic_score')}%")
        print(f"  ✅ [PASS] Kỹ năng Khớp (Matched): {matched}")
        print(f"  ⚠️ [PASS] Kỹ năng Thiếu (Missing): {missing}")
        print(f"  📝 [PASS] Nhận xét tự nhiên (Summary): \"{summary}\"")
    else:
        print(f"  ❌ [FAIL] Mong đợi 200 OK nhưng nhận mã {status_code}")

    # -----------------------------------------------------------------------
    print("\n" + "-" * 75)
    print("[TEST 4] Bắt lỗi khi thiếu trường bắt buộc (jd_data.title rỗng)...")
    invalid_match_payload = {
        "cv_data": {"skills": ["React"], "experience_years": 1},
        "jd_data": {"title": "", "required_skills": ["React"]},
    }
    status_code, resp_data = send_json_request(MATCH_CV_URL, invalid_match_payload)
    print(f"  👉 HTTP Status: {status_code}")
    print("  👉 Response Payload:")
    print(json.dumps(resp_data, indent=2, ensure_ascii=False))

    if status_code in (400, 422) and resp_data.get("success") is False:
        print(f"  ✅ [PASS] Bắt lỗi validation thiếu trường chính xác với mã {status_code}!")
    else:
        print(f"  ❌ [FAIL] Mong đợi 400 hoặc 422 nhưng nhận mã {status_code}")

    print("\n" + "=" * 75)
    print("🏁 HOÀN TẤT BỘ KIỂM THỬ MICROSERVICE.")
    print("=" * 75)


if __name__ == "__main__":
    run_tests()
