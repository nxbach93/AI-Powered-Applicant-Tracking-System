"""
Test trực tiếp async hàm analyze_cv trong app.api.endpoints bằng asyncio.
Kiểm tra luồng xử lý thực tế:
1. CV Web Dev 2 năm vs JD Kế toán 3 năm -> Điểm thấp <= 25%, Cảnh báo đầy đủ
2. CV Kế toán 5 năm vs JD Kế toán 3 năm -> Điểm cao >= 80%
"""

import sys
import os
import io
import asyncio

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

import docx
from starlette.datastructures import UploadFile
from app.api.endpoints import analyze_cv

def create_docx(name: str, title: str, exp_desc: str, skills_desc: str, edu_desc: str) -> bytes:
    doc = docx.Document()
    doc.add_heading(f"{name} - {title}", level=1)
    doc.add_paragraph("Email: candidate@example.com | Phone: 0987654321")
    doc.add_paragraph(f"HỌC VẤN: {edu_desc}")
    doc.add_paragraph(f"KINH NGHIỆM: {exp_desc}")
    doc.add_paragraph(f"KỸ NĂNG: {skills_desc}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

async def main():
    print("=" * 70)
    print("▶️ TEST ENDPOINT analyze_cv: CV Web Dev vs JD Kế toán 3 năm")
    print("=" * 70)
    
    cv_bytes = create_docx(
        name="Nguyễn Văn B",
        title="Frontend Web Developer",
        edu_desc="Đại học Bách Khoa Hà Nội, chuyên ngành Công nghệ thông tin",
        exp_desc="2 năm kinh nghiệm lập trình React, TypeScript, xây dựng website thương mại điện tử.",
        skills_desc="React, TypeScript, JavaScript, CSS, HTML, Git, Docker, Node.js",
    )
    
    jd_accounting = """Nhân viên Kế toán tổng hợp
Công ty TNHH Đầu tư & Phát triển ABC
Mô tả công việc: Lập báo cáo tài chính, quyết toán thuế, theo dõi dòng tiền, quản lý chứng từ sổ sách.
Yêu cầu:
- Tốt nghiệp Đại học chuyên ngành Kế toán, Kiểm toán.
- Tối thiểu 3 năm kinh nghiệm kế toán thực tế.
- Thành thạo MISA, Excel, nắm vững luật thuế."""

    upload_file = UploadFile(filename="cv_webdev.docx", file=io.BytesIO(cv_bytes))
    
    res = await analyze_cv(cv_file=upload_file, jd_text=jd_accounting)
    if hasattr(res, "body"):
        print("JSONResponse status:", res.status_code)
        print("JSONResponse body:", res.body.decode("utf-8"))
    print("Response:", res)
    print("  Match Score:", res.get("match_score"))
    print("  Score Breakdown:", res.get("score_breakdown"))
    print("  Domain Detected:", res.get("domain_detected"))
    print("  JD Domain Detected:", res.get("jd_domain_detected"))
    print("  Experience Gap:", res.get("experience_gap"))
    print("  Warnings:", res.get("warnings"))
    print("  Summary:", res.get("summary"))
    
    assert res.get("success") is True
    assert res.get("match_score") <= 25.0, f"Expected <= 25.0, got {res.get('match_score')}"
    assert len(res.get("warnings")) > 0
    print("✅ TEST analyze_cv trái ngành: PASS!")

    print("\n" + "=" * 70)
    print("▶️ TEST ENDPOINT analyze_cv: CV Kế toán 5 năm vs JD Kế toán 3 năm")
    print("=" * 70)
    
    cv_acc_bytes = create_docx(
        name="Lê Thị Lan",
        title="Kế toán trưởng",
        edu_desc="Đại học Kinh tế Quốc dân, Chuyên ngành Kế toán",
        exp_desc="5 năm kinh nghiệm kế toán tổng hợp, quyết toán thuế, lập báo cáo tài chính cho doanh nghiệp.",
        skills_desc="Kế toán tổng hợp, Kế toán thuế, Báo cáo tài chính, MISA, Excel, Kiểm toán",
    )
    
    upload_acc = UploadFile(filename="cv_ketoan.docx", file=io.BytesIO(cv_acc_bytes))
    res_acc = await analyze_cv(cv_file=upload_acc, jd_text=jd_accounting)
    print("Response:")
    print("  Success:", res_acc.get("success"))
    print("  Match Score:", res_acc.get("match_score"))
    print("  Score Breakdown:", res_acc.get("score_breakdown"))
    print("  Experience Gap:", res_acc.get("experience_gap"))
    print("  Warnings:", res_acc.get("warnings"))
    
    assert res_acc.get("success") is True
    assert res_acc.get("match_score") >= 80.0, f"Expected >= 80.0, got {res_acc.get('match_score')}"
    print("✅ TEST analyze_cv đúng ngành chuẩn kinh nghiệm: PASS!")

if __name__ == "__main__":
    asyncio.run(main())
